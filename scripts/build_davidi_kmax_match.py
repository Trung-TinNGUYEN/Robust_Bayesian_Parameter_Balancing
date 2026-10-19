"""Build data/external/davidi2016_kmax_matched.csv.

Maps the in vivo maximal apparent turnover numbers (kmax) of Davidi et al.
(2016, PNAS 113:3401, Dataset S1, doi 10.1073/pnas.1514240113) onto the
forward turnover-number rows (kcat+) of the four yeast-GEM pathway models used
in this repository (glycolysis, TCA cycle, pentose phosphate, shikimate).

Only derived per-reaction numbers are written.  The Davidi workbook itself is
not redistributed: download it from the article page and pass its path as the
first argument.

Inputs
  XLSX            pnas.1514240113.sd01.xlsx (PNAS Dataset S1), for example from
                  https://pmc.ncbi.nlm.nih.gov/articles/instance/4812741/bin/pnas.1514240113.sd01.xlsx
  --ijo1366       BiGG iJO1366 model in JSON (http://bigg.ucsd.edu/static/models/iJO1366.json),
                  used for gene rules, EC annotations and reaction directions
  --model-genes   model_genes.csv of github.com/dandanvidi/kvivo_max (MIT licence),
                  used for the protein molecular weights of the specific-activity conversion
  --s2            S1_Table.csv of supporting_tables/, used only to flag which
                  forward rows carry BRENDA kcat records (read-only)

Outputs
  --out           the derived table (default data/external/davidi2016_kmax_matched.csv)
  --full-out      optional working table with every intermediate column (Davidi
                  condition names, specific activities, S1 Table intervals).  It quotes
                  more of the source dataset, so write it outside the repository.

Rules
  * kmax is taken per polypeptide chain from the "kmax [1/s]" sheet wherever
    Davidi gives it (kmax_source = direct_1s_per_chain).
  * Otherwise it is converted from the "kmax [umol/mg/min]" sheet as
    kmax [1/s] = SA [umol/mg/min] x MW [kDa] / 60 (kmax_source =
    converted_from_specific_activity), with MW the chain mass for a single
    enzyme, the mean chain mass over isozymes for an isozyme reaction, the
    summed mass of one copy of each subunit for a complex, or a named isozyme
    where one dominates.  The rule used is in mw_rule.  This formula reproduces
    the directly given 1/s values to within about 5% for every single enzyme
    except PFK.
  * Directions: iJO1366 writes PGK, PGM, RPI and SUCOAS opposite to our models,
    so Davidi's "_reverse" entries are our forward direction.
  * The EC column of the SI kcat sheet is wrong for several rows, so EC matching
    uses the iJO1366 annotations.

Usage
  python3 scripts/build_davidi_kmax_match.py /path/to/pnas.1514240113.sd01.xlsx \\
      --ijo1366 /path/to/iJO1366.json --model-genes /path/to/model_genes.csv
Requires openpyxl (the isozyme class is read from the SI font colour).
"""
import argparse
import csv
import json
import math
import os
import re
import sys

import openpyxl

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_S2 = os.path.join(REPO, "supporting_tables", "S1_Table.csv")
DEFAULT_OUT = os.path.join(REPO, "data", "external", "davidi2016_kmax_matched.csv")

# Q-row layout (build_pathway_inputs.jl): kcat+ rows then kcat- rows, in pathway
# reaction order; the integer is the Q row of the first reaction's kcat+.
ORDER = {
    "glycolysis": (["r_0534", "r_0467", "r_0886", "r_0450", "r_1054", "r_0486",
                    "r_0892", "r_0893", "r_0366", "r_0962"], 51),
    "tca": (["r_0300", "r_0302", "r_0280", "r_0658", "r_0832", "r_0831", "r_0505",
             "r_1022", "r_1021", "r_0451", "r_0713"], 69),
    "pentose_phosphate": (["r_0466", "r_0091", "r_0889", "r_0982", "r_0984",
                           "r_1049", "r_1050", "r_1048"], 43),
    "shikimate": (["r_0042", "r_0040", "r_0039", "r_0996", "r_0997", "r_0065",
                   "r_0279"], 42),
}

# (pathway, yeast reaction, name, our EC, kcat+ reference substrate,
#  iJO1366 reaction, Davidi id for OUR forward direction (None = absent),
#  match quality, note, MW rule for the specific-activity conversion)
SPEC = [
    ("glycolysis", "r_0534", "hexokinase", "2.7.1.1", "D-glucose", "HEX1", None, "none",
     "E. coli has no hexokinase: glucose enters via PTS; HEX1 (glk, glucokinase) carries no pFBA flux and has no kmax.", None),
    ("glycolysis", "r_0467", "glucose-6-phosphate isomerase", "5.3.1.9", "D-glucose 6-phosphate", "PGI", "PGI", "exact", "", "chain"),
    ("glycolysis", "r_0886", "phosphofructokinase", "2.7.1.11", "D-fructose 6-phosphate", "PFK", "PFK", "exact",
     "Davidi assigned PFK to pfkA only (pfkB minor). The SI records 1 catalytic site per pfkA tetramer, so the per-active-site kmax is 4x the per-chain value; the specific-activity maximum (about 127 1/s after conversion) disagrees with the per-chain 278 1/s.", "chain"),
    ("glycolysis", "r_0450", "fructose-bisphosphate aldolase", "4.1.2.13", "D-fructose 1,6-bisphosphate", "FBA", "FBA", "exact",
     "Isozymes fbaA (class II), fbaB (class I), ydjI: the specific activity is a mass-weighted mean over isozymes, so it underestimates the fastest isozyme; approximate conversion.", "isozyme"),
    ("glycolysis", "r_1054", "triose-phosphate isomerase", "5.3.1.1", "dihydroxyacetone phosphate", "TPI", "TPI", "exact",
     "SI kcat sheet EC column wrongly shows 4.1.1.31 for TPI.", "chain"),
    ("glycolysis", "r_0486", "glyceraldehyde-3-phosphate dehydrogenase", "1.2.1.12", "D-glyceraldehyde 3-phosphate", "GAPD", "GAPD", "exact", "", "chain"),
    ("glycolysis", "r_0892", "phosphoglycerate kinase", "2.7.2.3", "3-phospho-D-glyceroyl phosphate", "PGK", "PGK_reverse", "exact",
     "iJO1366 writes PGK in the gluconeogenic direction; PGK_reverse (13dpg + adp -> 3pg + atp) is our forward.", "chain"),
    ("glycolysis", "r_0893", "phosphoglycerate mutase", "5.4.2.11", "3-phospho-D-glycerate", "PGM", "PGM_reverse", "exact",
     "iJO1366 writes PGM as 2pg -> 3pg; PGM_reverse is our forward. Isozymes gpmA, gpmM, ytjC; MW taken as gpmA (assumed dominant); approximate conversion.", "gpmA"),
    ("glycolysis", "r_0366", "enolase", "4.2.1.11", "2-phospho-D-glycerate", "ENO", "ENO", "exact", "", "chain"),
    ("glycolysis", "r_0962", "pyruvate kinase", "2.7.1.40", "phosphoenolpyruvate", "PYK", None, "none",
     "PYK (pykF, pykA) is absent from every Davidi kapp and kmax sheet.", None),
    ("tca", "r_0300", "citrate synthase", "2.3.3.1", "acetyl-CoA", "CS", "CS", "exact",
     "Yeast r_0300 is mitochondrial; E. coli gltA is a type II hexameric citrate synthase.", "chain"),
    ("tca", "r_0302", "aconitase step 1 (citrate -> cis-aconitate)", "4.2.1.3", "citrate", "ACONTa", "ACONTa", "exact",
     "Isozymes acnA, acnB; mass-weighted specific activity; approximate conversion. Davidi reports the same specific activity for ACONTa and ACONTb.", "isozyme"),
    ("tca", "r_0280", "aconitase step 2 (cis-aconitate -> isocitrate)", "4.2.1.3", "cis-aconitate", "ACONTb", "ACONTb", "exact",
     "No BRENDA kcat records on this row: cis-aconitate kcat records land on kcat-(r_0302) by the last-write-wins lookup. Isozymes; approximate conversion.", "isozyme"),
    ("tca", "r_0658", "isocitrate dehydrogenase (NAD+)", "1.1.1.41", "isocitrate", "ICDHyr", "ICDHyr", "cofactor_differs",
     "Yeast r_0658 is NAD+-dependent (EC 1.1.1.41); E. coli icd is NADP+-dependent (EC 1.1.1.42). Reaction-level match only.", "chain"),
    ("tca", "r_0832", "2-oxoglutarate dehydrogenase E1", "1.2.4.2", "2-oxoglutarate", "AKGDH", "AKGDH", "lumped_complex",
     "iJO1366 AKGDH is the whole OGDH complex (sucA E1, sucB E2, lpd E3) whereas r_0832 is the E1 half-reaction; conversion per one sucA+sucB+lpd set is nominal.", "complex"),
    ("tca", "r_0831", "2-oxoglutarate dehydrogenase E2", "2.3.1.61", "succinyldihydrolipoamide", "AKGDH", "AKGDH", "lumped_complex",
     "Lumped AKGDH (see r_0832). No BRENDA kcat records on this row.", "complex"),
    ("tca", "r_0505", "dihydrolipoyl dehydrogenase (E3)", "1.8.1.4", "DL-dihydrolipoamide", "AKGDH", "AKGDH", "lumped_complex",
     "Lumped AKGDH (see r_0832); lpd is shared with PDH and glycine cleavage, so AKGDH flux understates E3 turnover. No BRENDA kcat records on this row.", "complex"),
    ("tca", "r_1022", "succinyl-CoA ligase (ADP-forming)", "6.2.1.5", "succinyl-CoA", "SUCOAS", "SUCOAS_reverse", "exact",
     "iJO1366 writes SUCOAS as succ -> succoa; SUCOAS_reverse is our forward. sucC2sucD2 heterotetramer; conversion per one sucC+sucD pair.", "complex"),
    ("tca", "r_1021", "succinate dehydrogenase", "1.3.5.1", "succinate", "SUCDi", "SUCDi", "exact",
     "Electron acceptor differs (ubiquinone-6 in yeast, ubiquinone-8 in E. coli); BiGG gives no EC (1.3.5.1 by function). sdhABCD complex; conversion per one protomer.", "complex"),
    ("tca", "r_0451", "fumarase", "4.2.1.2", "fumarate", "FUM", "FUM", "exact",
     "Isozymes fumA, fumB (class I) and fumC (class II); mass-weighted specific activity; approximate conversion.", "isozyme"),
    ("tca", "r_0713", "malate dehydrogenase", "1.1.1.37", "L-malate", "MDH", "MDH", "exact",
     "SI kcat sheet EC column wrongly shows 3.5.1.16 for MDH. MDH runs close to equilibrium in vivo (Davidi kcat/kmax = 144).", "chain"),
    ("pentose_phosphate", "r_0466", "glucose-6-phosphate dehydrogenase", "1.1.1.49", "D-glucose 6-phosphate", "G6PDH2r", "G6PDH2r", "exact", "", "chain"),
    ("pentose_phosphate", "r_0091", "6-phosphogluconolactonase", "3.1.1.31", "6-phospho-D-glucono-1,5-lactone", "PGL", "PGL", "exact",
     "No BRENDA kcat records for 3.1.1.31, so the row carries no records. SI kcat sheet EC column wrongly shows 5.3.1.1 for PGL.", "chain"),
    ("pentose_phosphate", "r_0889", "phosphogluconate dehydrogenase", "1.1.1.44", "6-phosphogluconate", "GND", "GND", "exact", "", "chain"),
    ("pentose_phosphate", "r_0982", "ribose-5-phosphate isomerase", "5.3.1.6", "D-ribulose 5-phosphate", "RPI", "RPI_reverse", "exact",
     "iJO1366 writes RPI as r5p -> ru5p; RPI_reverse is our forward. Isozymes rpiA, rpiB; MW taken as rpiA (assumption); approximate conversion.", "rpiA"),
    ("pentose_phosphate", "r_0984", "ribulose-5-phosphate 3-epimerase", "5.1.3.1", "D-ribulose 5-phosphate", "RPE", "RPE", "exact",
     "Isozymes rpe, sgcE (Davidi's code assigns RPE to rpe only); MW taken as rpe; approximate conversion.", "rpe"),
    ("pentose_phosphate", "r_1049", "transketolase 1", "2.2.1.1", "D-xylulose 5-phosphate", "TKT1", "TKT1", "exact",
     "Isozymes tktA, tktB; mass-weighted specific activity; approximate conversion.", "isozyme"),
    ("pentose_phosphate", "r_1050", "transketolase 2", "2.2.1.1", "D-erythrose 4-phosphate", "TKT2", "TKT2", "exact",
     "Isozymes tktA, tktB; approximate conversion.", "isozyme"),
    ("pentose_phosphate", "r_1048", "transaldolase", "2.2.1.2", "glyceraldehyde 3-phosphate", "TALA", None, "none",
     "TALA (talA, talB) is absent from every Davidi kapp and kmax sheet.", None),
    ("shikimate", "r_0042", "DAHP synthase", "2.5.1.54", "phosphoenolpyruvate", "DDPA", "DDPA", "exact",
     "Isozymes aroF, aroG, aroH (feedback-sensitive to Tyr, Phe, Trp); approximate conversion.", "isozyme"),
    ("shikimate", "r_0040", "3-dehydroquinate synthase", "4.2.3.4", "3-deoxy-D-arabino-hept-2-ulosonate 7-phosphate", "DHQS", "DHQS", "exact",
     "Yeast enzyme is the pentafunctional ARO1; E. coli aroB is monofunctional. SI EC column wrongly shows 5.3.1.16.", "chain"),
    ("shikimate", "r_0039", "3-dehydroquinate dehydratase", "4.2.1.10", "3-dehydroquinate", "DHQTi", "DHQTi", "exact",
     "ARO1 (yeast) against monofunctional type I aroD (E. coli).", "chain"),
    ("shikimate", "r_0996", "shikimate dehydrogenase", "1.1.1.25", "3-dehydroshikimate", "SHK3Dr", "SHK3Dr", "exact",
     "Same (reductive) direction in both models. Isozymes aroE, ydiB; MW taken as aroE (assumption); approximate conversion.", "aroE"),
    ("shikimate", "r_0997", "shikimate kinase", "2.7.1.71", "shikimate", "SHKK", "SHKK", "exact",
     "Isozymes aroK, aroL; approximate conversion.", "isozyme"),
    ("shikimate", "r_0065", "3-phosphoshikimate 1-carboxyvinyltransferase", "2.5.1.19", "3-phosphoshikimate", "PSCVT", "PSCVT", "exact", "", "chain"),
    ("shikimate", "r_0279", "chorismate synthase", "4.2.3.5", "5-enolpyruvylshikimate 3-phosphate", "CHORS", "CHORS", "exact",
     "SI EC column wrongly shows 2.1.2.2 for CHORS.", "chain"),
]

NAMED_GENE = {"gpmA": "b0755", "rpiA": "b2914", "rpe": "b3386", "aroE": "b3281"}

COLUMNS = [
    "pathway", "reaction_id", "reaction_name", "direction", "Q_row", "ec_code",
    "reference_substrate", "has_kcat_records", "iJO1366_reaction", "iJO1366_genes",
    "davidi_reaction_id", "match_quality", "match_by_reaction", "match_by_ec",
    "kmax_1s", "kmax_source", "mw_used_kDa", "mw_rule", "kmax_per_active_site_1s",
    "isozymes_in_iJO1366", "protein_complex", "davidi_si_enzyme_class", "notes",
]


def fmt(x, nd=4):
    if x is None or x == "":
        return ""
    if isinstance(x, str):
        return x
    return "{:.{}g}".format(x, nd)


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("xlsx", help="Davidi et al. 2016 Dataset S1 (pnas.1514240113.sd01.xlsx)")
    ap.add_argument("--ijo1366", required=True, help="BiGG iJO1366 JSON model")
    ap.add_argument("--model-genes", required=True, help="kvivo_max model_genes.csv (tab-separated)")
    ap.add_argument("--s2", default=DEFAULT_S2, help="S1_Table.csv of supporting_tables/ (read-only)")
    ap.add_argument("--out", default=DEFAULT_OUT)
    ap.add_argument("--full-out", default=None)
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.xlsx, data_only=True)

    def sheet_rows(name, header_row):
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        hdr = rows[header_row - 1]
        out, colour = {}, {}
        for i, r in enumerate(rows[header_row:], start=header_row + 1):
            if r[0] is None:
                continue
            out.setdefault(r[0], dict(zip(hdr, r)))
            c = ws.cell(row=i, column=1).font.color
            colour[r[0]] = c.rgb if (c is not None and c.type == "rgb") else None
        return out, colour

    kcat_s, _ = sheet_rows("kcat 1s", 3)
    kmax_s, _ = sheet_rows("kmax 1s", 3)
    kvk_s, _ = sheet_rows("kcat vs. kmax", 3)
    sa_s, sa_col = sheet_rows("kmax umolmgmin", 4)

    with open(args.ijo1366) as f:
        model = json.load(f)
    R = {r["id"]: r for r in model["reactions"]}
    gname = {g["id"]: g.get("name") for g in model["genes"]}
    MW = {}
    with open(args.model_genes) as f:
        for row in csv.DictReader(f, delimiter="\t"):
            if row["MW_Da"]:
                MW[row["bnumber"]] = float(row["MW_Da"]) / 1000.0

    s2 = {}
    with open(args.s2) as f:
        for r in csv.DictReader(f):
            if r["kind"] == "kcat":
                s2[(r["pathway"], int(r["Q_row"]))] = r

    def gpr_genes(rid):
        return re.findall(r"b\d{4}", R[rid]["gene_reaction_rule"])

    def rxn_string(rid, reverse):
        s = R[rid]["metabolites"]
        lhs = [k[:-2] for k, v in s.items() if v < 0]
        rhs = [k[:-2] for k, v in s.items() if v > 0]
        if reverse:
            lhs, rhs = rhs, lhs
        return " + ".join(lhs) + " -> " + " + ".join(rhs)

    rows, full = [], []
    for (pw, yr, yname, ec, sub, ijo, did, quality, note, mwrule) in SPEC:
        rx, first = ORDER[pw]
        q = first + rx.index(yr)
        srow = s2.get((pw, q))
        if srow is not None:
            assert srow["ec_code"] == ec and srow["substrate"].lower() == sub.lower(), (pw, yr, q)
        rule = R[ijo]["gene_reaction_rule"]
        genes = gpr_genes(ijo)
        ecs = R[ijo].get("annotation", {}).get("ec-code", []) or []
        row = dict.fromkeys(COLUMNS, "")
        row.update(pathway=pw, reaction_id=yr, reaction_name=yname, direction="forward",
                   Q_row=q, ec_code=ec, reference_substrate=sub,
                   has_kcat_records="true" if srow is not None else "false",
                   iJO1366_reaction=ijo,
                   iJO1366_genes=re.sub(r"b\d{4}", lambda m: gname.get(m.group(0)) or m.group(0), rule),
                   match_quality=quality,
                   isozymes_in_iJO1366="true" if " or " in rule else "false",
                   protein_complex="true" if " and " in rule else "false",
                   notes=note)
        if ec in ecs:
            row["match_by_ec"] = "yes"
        elif ijo == "SUCDi":
            row["match_by_ec"] = "yes (by function; BiGG lists no EC)"
        else:
            row["match_by_ec"] = "no (iJO1366 EC " + ";".join(ecs) + ")"
        extra = {}
        if did is None:
            row["match_by_reaction"] = "no (absent from the Davidi kmax sheets)"
            rows.append(row)
            full.append(dict(row, **extra))
            continue
        rev = did.endswith("_reverse")
        row["davidi_reaction_id"] = did
        row["match_by_reaction"] = {
            "exact": "yes" + (" (Davidi _reverse is our forward)" if rev else ""),
            "cofactor_differs": "yes, cofactor differs (NADP+ in E. coli, NAD+ in yeast)",
            "lumped_complex": "partial (lumped OGDH complex)",
        }[quality]
        row["davidi_si_enzyme_class"] = {"FF333399": "single enzyme", "FF993300": "isozymes"}.get(sa_col.get(did), "")
        extra["davidi_direction_string"] = rxn_string(ijo, rev)
        best, src = None, ""
        k = kmax_s.get(did)
        if k:
            best, src = k["kmax per polypeptide chain [s-1]"], "direct_1s_per_chain"
            extra["kmax_condition"] = k["condition"]
            kv = kvk_s.get(did)
            if kv:
                row["kmax_per_active_site_1s"] = fmt(kv["kmax [s-1]"])
            elif did in kcat_s:
                kc = kcat_s[did]
                row["kmax_per_active_site_1s"] = fmt(best * kc["polypeptides per complex"] /
                                                     kc["catalytic sites per complex"])
        sa = sa_s.get(did)
        if sa:
            if mwrule == "chain":
                mw, mwtxt = MW[genes[0]], "single chain (" + gname[genes[0]] + ")"
            elif mwrule == "isozyme":
                ms = [MW[g] for g in genes]
                mw = sum(ms) / len(ms)
                mwtxt = "mean chain mass over isozymes ({:.1f}-{:.1f} kDa)".format(min(ms), max(ms))
            elif mwrule == "complex":
                mw = sum(MW[g] for g in genes)
                mwtxt = "sum of one copy of each subunit"
            else:
                mw, mwtxt = MW[NAMED_GENE[mwrule]], "named isozyme (" + mwrule + ")"
            conv = sa["kmax [umol/mg/min]"] * mw / 60.0
            extra["kmax_specific_activity_umol_mg_min"] = sa["kmax [umol/mg/min]"]
            extra["SA_condition"] = sa["condition"]
            extra["kmax_from_SA_1s"] = conv
            if best is None:
                best, src = conv, "converted_from_specific_activity"
                row["mw_used_kDa"] = "{:.2f}".format(mw)
                row["mw_rule"] = mwtxt
        row["kmax_1s"] = fmt(best)
        row["kmax_source"] = src
        kc = kcat_s.get(did)
        if kc:
            extra["kcat_vitro_ecoli_davidi_1s"] = kc["kcat per active site [1/s]"]
            extra["kcat_vitro_SI_ec_column"] = kc["EC"]
        if srow is not None:
            for c in ("n_records", "n_organisms", "consensus_median", "consensus_lo",
                      "consensus_hi", "neworganism_lo", "neworganism_hi"):
                extra["S2_" + c] = srow[c]
        rows.append(row)
        full.append(dict(row, **extra))

    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    with open(args.out, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS)
        w.writeheader()
        w.writerows(rows)
    print("wrote", args.out, len(rows), "rows")
    if args.full_out:
        keys = COLUMNS + sorted({k for r in full for k in r} - set(COLUMNS))
        with open(args.full_out, "w", newline="") as f:
            w = csv.DictWriter(f, fieldnames=keys)
            w.writeheader()
            w.writerows(full)
        print("wrote", args.full_out)

    matched = [r for r in rows if r["kmax_1s"]]
    rec = [r for r in matched if r["has_kcat_records"] == "true"]
    direct = [r for r in rec if r["kmax_source"] == "direct_1s_per_chain"]
    print("forward rows: {}; with kmax: {}; with kmax and kcat records: {} ({} given directly in 1/s)".format(
        len(rows), len(matched), len(rec), len(direct)))
    for pw in ORDER:
        sub = [r for r in rec if r["pathway"] == pw]
        print("  {:18s} matched rows with records {:2d}, direct {:2d}, converted {:2d}".format(
            pw, len(sub), sum(r["kmax_source"].startswith("direct") for r in sub),
            sum(r["kmax_source"].startswith("converted") for r in sub)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
